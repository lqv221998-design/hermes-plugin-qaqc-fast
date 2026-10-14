"""Hybrid QAQC behavioral tests. All business data is synthetic."""
import importlib.util
import json
from pathlib import Path
import sys
import os
import tempfile
import unittest
from typing import Any, cast

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('qaqc_expansion_test', ROOT/'__init__.py', submodule_search_locations=[str(ROOT)])
assert spec is not None and spec.loader is not None
plugin = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = plugin
spec.loader.exec_module(plugin)
URL = 'https://qaqc.newtecons.vn/#/AlbumThiCong/1/synthetic'

def snapshot(body='', project='Fixture Project', discipline=''):
    return ('uid=1_0 RootWebArea "Fixture" url="'+URL+'"\n'
            '  uid=1_1 StaticText "Fixture User"\n'
            '  uid=1_2 link "QAQC" url="unsafe:javascript:;"\n'
            f'  uid=1_3 StaticText "{project}"\n'
            '  uid=1_4 StaticText "XÂY DỰNG"\n'
            '  uid=1_5 StaticText "CƠ ĐIỆN"\n'
            '  uid=1_6 StaticText "HỆ THỐNG"\n'
            '  uid=1_7 StaticText "PHÁP CHẾ"\n'+discipline+body)

class State:
    def __init__(self, root):
        self.data_dir = Path(root)
        self.values = {}
    def get(self,key,default=None):
        return self.values.get(key,default)
    def set(self,key,value):
        self.values[key] = value

class Host:
    def __init__(self, root, outputs):
        self.state = State(root)
        self.outputs = iter(outputs)
        self.calls = []
    def get_config(self,key,default=None):
        return 'manual' if key=='context_policy' else 'Fixture User' if key=='expected_user_display_name' else default
    def call_mcp(self,server,tool,arguments,timeout=15):
        self.calls.append((tool,arguments))
        text = next(self.outputs)
        if tool=='take_snapshot' and 'filePath' in arguments:
            Path(arguments['filePath']).write_text(text,encoding='utf-8')
            return {'ok':True,'result':'Snapshot saved to '+arguments['filePath']}
        return {'ok':True,'result':text}

def extended(host,args) -> Any:
    fn = cast(Any,getattr(plugin,'run_extended',None))
    assert callable(fn), 'Missing extended native executor'
    return cast(Any,fn(host,args,session_id='fixture-session'))

class Expansion(unittest.TestCase):
    def test_setup_skill_is_registered_without_browser_or_install_calls(self):
        class Context:
            def __init__(self):
                self.skills = {}
            def register_tool(self,**kwargs):
                pass
            def register_hook(self,*args):
                pass
            def register_skill(self,name,path,**kwargs):
                self.skills[name] = path
        ctx = Context()
        plugin.register(ctx)
        self.assertIn('setup',ctx.skills)
        self.assertTrue(ctx.skills['setup'].is_file())

    def test_native_input_recaptures_exact_window_and_blocks_changed_pixels(self):
        import hashlib,time
        with tempfile.TemporaryDirectory() as root:
            before,after = Path(root)/'before.png',Path(root)/'after.png'
            before.write_bytes(b'synthetic reviewed frame')
            after.write_bytes(b'synthetic changed frame')
            digest = hashlib.sha256(before.read_bytes()).hexdigest()
            record = 'QAQC-PLUGIN-TEST-FIXTURE-CUA-STALE'
            host = Host(root,[])
            old = host.get_config
            host.get_config = lambda key,default=None: 'Fixture Project' if key=='write_project' else old(key,default)
            host.state.set('native_handoff',{'token':'fixture','session_id':'fixture-session','project':'Fixture Project','user':'Fixture User','source_url':URL,'record_id':record,'pid':1,'window_id':2,'captured_at':time.time(),'sha256':digest,'screenshot_path':str(before),'elements':{'1':'Canvas','2':'Canvas'},'consumed':False})
            key = 'owned_'+hashlib.sha256(('Fixture Project\0'+record).encode()).hexdigest()
            host.state.set(key,{'user':'Fixture User','source_url':URL,'status':'created_verified'})
            inputs = []
            def dispatch(name,args,**kwargs):
                if args['action']=='capture':
                    return json.dumps({'screenshot_path':str(after),'summary':'changed fixture frame'})
                inputs.append(args)
                return json.dumps({'ok':True,'verified':True})
            cast(Any,host).dispatch_tool = dispatch
            result = extended(host,dict(action='native_step',page_id=1,expected_project='Fixture Project',record_id=record,operation_id='stale-native',handoff_token='fixture',reviewed_capture_sha256=digest,purpose='canvas',native_action={'action':'drag','from_element':1,'to_element':2}))
            self.assertFalse(result['ok'],result)
            self.assertEqual(inputs,[])
            self.assertEqual(result['status'],'native_live_capture_changed_review_required')

    def test_transport_refusal_invalidates_native_handoff(self):
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[])
            host.state.set('native_handoff',{'token':'old','project':'Fixture Project'})
            host.call_mcp = lambda *args,**kw: {'ok':False,'error':'fixture refusal'}
            result = extended(host,dict(action='read_view',page_id=1,expected_project='Fixture Project'))
            self.assertFalse(result['ok'])
            self.assertTrue(cast(Any,host.state.get('native_handoff')).get('invalidated'))

    def test_native_capture_refreshes_owned_dialog_without_dom_fallback(self):
        import time
        with tempfile.TemporaryDirectory() as root:
            image = Path(root)/'dialog.png'
            image.write_bytes(b'synthetic native dialog')
            host = Host(root,[])
            host.state.set('native_handoff',{'token':'old-token','session_id':'fixture-session','project':'Fixture Project','user':'Fixture User','source_url':URL,'source_verification':'native_header_user_project_checked','pid':11,'window_id':22,'captured_at':time.time(),'consumed':True})
            cast(Any,host).dispatch_tool = lambda name,args,**kw: json.dumps({'screenshot_path':str(image),'text_summary':'Save As native dialog','_multimodal':True,'content':[{'type':'text','text':'Synthetic dialog'}]})
            result = extended(host,dict(action='native_capture',page_id=1,expected_project='Fixture Project',pid=11,window_id=22,handoff_token='old-token'))
            self.assertTrue(result['ok'],result)
            self.assertEqual(host.calls,[])
            self.assertNotEqual(result['handoff']['token'],'old-token')
            self.assertTrue(result['handoff']['visual_review_required'])

    def test_failed_reload_never_proves_persistent_write(self):
        import hashlib
        record = 'QAQC-PLUGIN-TEST-FIXTURE-RELOAD'
        target = 'https://qaqc.newtecons.vn/#/editForm/fixture/42/1/project'
        before = snapshot('  uid=8_0 textbox "Mã hồ sơ" value="'+record+'"\n  uid=8_1 textbox "Nội dung" value="Old"\n  uid=8_2 button "Lưu"\n  uid=8_3 combobox "Trạng thái" value="Draft"\n').replace(URL,target)
        after = before.replace('value="Old"','value="New"')
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[before,after,after,'Unable to reload the selected page: navigation cancelled.',after])
            old = host.get_config
            host.get_config = lambda key,default=None: 'Fixture Project' if key=='write_project' else old(key,default)
            key = 'owned_'+hashlib.sha256(('Fixture Project\0'+record).encode()).hexdigest()
            host.state.set(key,{'project':'Fixture Project','user':'Fixture User','status':'created_verified'})
            result = extended(host,dict(action='write_record',page_id=1,expected_project='Fixture Project',record_id=record,operation_id='failed-reload',intent='edit',fields={'Nội dung':'New'},submit_label='Lưu',expected_status='Draft'))
            self.assertFalse(result['ok'],result)
            self.assertEqual(result['status'],'reload_not_verified')
            self.assertEqual(len(host.calls),4)

    def test_public_read_list_uses_new_readiness_verification(self):
        class RegisteredHost(Host):
            def __init__(self,root,outputs):
                super().__init__(root,outputs)
                self.tools = {}
            def register_tool(self,**kwargs):
                self.tools[kwargs['name']] = kwargs
            def register_hook(self,*args):
                pass
            def register_skill(self,*args,**kwargs):
                pass
        text = snapshot('  uid=8_0 heading "QUẢN LÝ HÌNH ẢNH THI CÔNG"\n  uid=8_1 table busy=false\n    uid=8_2 row\n      uid=8_3 cell "R001"\n')
        with tempfile.TemporaryDirectory() as root:
            host = RegisteredHost(root,[text,text])
            plugin.register(host)
            result = json.loads(host.tools['qaqc_fast_session']['handler'](dict(action='read_list',module='photos',page_id=1,expected_project='Fixture Project'),session_id='fixture-session'))
            self.assertTrue(result.get('data_ready_verified'),result)

    def test_benchmark_records_latency_and_output_stability(self):
        text = snapshot('  uid=8_0 table busy=false\n    uid=8_1 row\n      uid=8_2 cell "R001"\n')
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[text]*5)
            result = extended(host,dict(action='benchmark',page_id=1,expected_project='Fixture Project',iterations=2))
            self.assertTrue(result['ok'],result)
            data = json.loads(Path(result['artifact']['path']).read_text(encoding='utf-8'))
            self.assertEqual(data['samples'],2)
            self.assertTrue(data['output_stable'])
            self.assertFalse(data['includes_model_and_chat'])
            self.assertEqual(result['mcp_calls'],5,'preflight plus two fresh snapshots for each sample')

    def test_readiness_accepts_stable_record_counter_but_not_loading(self):
        loaded = plugin.parse_snapshot(snapshot('  uid=8_0 StaticText "Total: 1"\n  uid=8_1 table\n    uid=8_2 row\n      uid=8_3 cell "R001"\n'))
        fn = cast(Any,plugin.readiness)
        self.assertTrue(fn(loaded,loaded)['data_ready_verified'])
        loading = plugin.parse_snapshot(snapshot('  uid=7_0 StaticText "Đang tải dữ liệu..."\n  uid=8_0 StaticText "Không có dữ liệu"\n'))
        self.assertFalse(fn(loading,loading)['data_ready_verified'])

    def test_checkpoint_reports_uncertain_state_without_replay(self):
        import hashlib
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[])
            digest = hashlib.sha256(b'fixture-pending').hexdigest()
            host.state.set('scope_'+digest,{'project':'Fixture Project','user':'Fixture User'})
            host.state.set('operation_'+digest,{'status':'inflight_uncertain_until_verified'})
            result = extended(host,dict(action='checkpoint',page_id=1,expected_project='Fixture Project',operation_id='fixture-pending'))
            self.assertTrue(result['ok'],result)
            self.assertEqual(result['operation']['status'],'inflight_uncertain_until_verified')
            self.assertFalse(result['automatic_resume_allowed'])
            self.assertEqual(host.calls,[])

    def test_restore_executor_refuses_ungranted_lifecycle_action(self):
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[])
            result = extended(host,dict(action='restore_browser',page_id=1,expected_project='Fixture Project',operation_id='restore-unbound'))
            self.assertEqual(result['status'],'browser_restore_not_authorized',result)
            self.assertEqual(host.calls,[])

    def test_recovery_launches_only_missing_exact_owner(self):
        from http.server import HTTPServer,BaseHTTPRequestHandler
        import threading
        fn = cast(Any,getattr(plugin,'recover_owned_browser',None))
        self.assertTrue(callable(fn),'Missing controlled browser recovery')
        assert callable(fn)
        with tempfile.TemporaryDirectory() as root:
            profile = Path(root)/'recipient-profile'
            profile.mkdir()
            exe = Path(root)/'chrome.exe'
            exe.write_bytes(b'synthetic executable; never executed')
            class Handler(BaseHTTPRequestHandler):
                def do_GET(self):
                    body = json.dumps({'Browser':'Chrome/150.0.0.0'}).encode()
                    self.send_response(200)
                    self.end_headers()
                    self.wfile.write(body)
                def log_message(self,*args):
                    pass
            server = HTTPServer(('127.0.0.1',0),Handler)
            port = server.server_address[1]
            server.server_close()
            owners,calls,servers = [],[],[]
            def launch(argv):
                calls.append(argv)
                running = HTTPServer(('127.0.0.1',port),Handler)
                servers.append(running)
                threading.Thread(target=running.serve_forever,daemon=True).start()
                owners.append({'Pid':42,'DebugPort':port,'UserDataDir':str(profile),'Executable':str(exe)})
            try:
                value = cast(Any,fn(exe,profile,port,launch_fn=launch,owner_probe=lambda:owners))
                self.assertTrue(value['ok'],value)
                self.assertTrue(value['restored'])
                self.assertEqual(len(calls),1)
                second = cast(Any,fn(exe,profile,port,launch_fn=launch,owner_probe=lambda:owners))
                self.assertFalse(second['restored'])
                self.assertEqual(len(calls),1)
            finally:
                for active in servers:
                    active.shutdown()
                    active.server_close()

    def test_selected_photo_uses_host_vision_without_becoming_normative(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'selected-photo.png'
            Image.new('RGB',(100,100),'green').save(path)
            host = Host(root,[snapshot()])
            cast(Any,host).dispatch_tool = lambda name,args,**kw: json.dumps({'success':True,'analysis':'Synthetic green image'})
            result = extended(host,dict(action='analyze_image',page_id=1,expected_project='Fixture Project',operation_id='vision-fixture',local_file=str(path),image_question='Describe the selected image; do not infer technical compliance.'))
            self.assertTrue(result['ok'],result)
            self.assertFalse(result['normative'])
            self.assertEqual(result['review_status'],'unreviewed')

    def test_ocr_executor_requires_explicit_local_service_binding(self):
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[snapshot()])
            result = extended(host,dict(action='ocr_document',page_id=1,expected_project='Fixture Project',operation_id='ocr-fixture-unbound',local_file=str(Path(root)/'selected.pdf')))
            self.assertEqual(result['status'],'explicit_local_mineru_binding_required',result)

    @unittest.skipUnless(os.environ.get('QAQC_MINERU_TEST_URL'),'Explicit local MinerU integration only')
    def test_mineru_extracts_real_selected_synthetic_image(self):
        from PIL import Image, ImageDraw, ImageFont
        fn = cast(Any,getattr(plugin,'mineru_parse_file',None))
        self.assertTrue(callable(fn),'Missing MinerU adapter')
        assert callable(fn)
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'synthetic-scan.png'
            image = Image.new('RGB',(1200,500),'white')
            draw = ImageDraw.Draw(image)
            font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf',48)
            draw.text((50,100),'QAQC SYNTHETIC OCR TEST',font=font,fill='black')
            draw.text((50,200),'Record R001. No company data.',font=font,fill='black')
            image.save(path)
            value = cast(Any,fn(path,os.environ['QAQC_MINERU_TEST_URL']))
            self.assertTrue(value['ok'],value)
            self.assertIn('QAQC',json.dumps(value['raw_response'],ensure_ascii=False).upper())

    def test_delete_verifies_record_absence_in_exact_return_list(self):
        import hashlib
        record = 'QAQC-PLUGIN-TEST-FIXTURE-DELETE'
        target = 'https://qaqc.newtecons.vn/#/editForm/fixture/42/1/project'
        before = snapshot('  uid=8_0 textbox "Mã hồ sơ" value="'+record+'"\n  uid=8_1 button "Xóa"\n  uid=8_2 link "Back to list" url="'+URL+'"\n').replace(URL,target)
        before = before.replace('link "Back to list" url="'+target+'"','link "Back to list" url="'+URL+'"')
        empty = snapshot('  uid=9_0 StaticText "Total: 0"\n  uid=9_1 StaticText "Không có dữ liệu"\n  uid=9_2 button "Next" disabled=true\n')
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[before,empty,'Successfully reloaded the page.',empty,empty])
            old = host.get_config
            host.get_config = lambda key,default=None: 'Fixture Project' if key=='write_project' else old(key,default)
            key = 'owned_'+hashlib.sha256(('Fixture Project\0'+record).encode()).hexdigest()
            host.state.set(key,{'project':'Fixture Project','user':'Fixture User','status':'created_verified','source_url':target})
            result = extended(host,dict(action='delete_record',page_id=1,expected_project='Fixture Project',record_id=record,operation_id='delete-fixture',control_label='Xóa',return_list_label='Back to list',next_label='Next',key_column=0,max_pages=5))
            self.assertTrue(result['ok'],result)
            self.assertEqual(result['status'],'record_absence_verified_in_return_list')
            self.assertEqual(host.state.get(key)['status'],'deleted_verified')

    def test_large_snapshot_is_read_losslessly_from_native_file(self):
        body = '  uid=8_0 table busy=false\n'
        body += ''.join(f'    uid=9_{i} row\n      uid=10_{i} cell "R{i}"\n' for i in range(1600))
        text = snapshot(body)
        self.assertGreater(len(text),65000)
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[text,text])
            result = extended(host,dict(action='read_view',page_id=1,expected_project='Fixture Project'))
            self.assertTrue(result['ok'],result)
            self.assertEqual(result['rows'],1600)

    def test_registered_tool_exposes_extended_actions_with_approval(self):
        class Registration:
            def __init__(self):
                self.tools,self.hooks = {},{}
            def register_tool(self,**kwargs):
                self.tools[kwargs['name']] = kwargs
            def register_hook(self,name,callback):
                self.hooks[name] = callback
            def register_skill(self,*args,**kwargs):
                pass
        ctx = Registration()
        plugin.register(ctx)
        tool = ctx.tools['qaqc_fast_session']
        self.assertIn('read_all',tool['schema']['parameters']['properties']['action']['enum'])
        verdict = ctx.hooks['pre_tool_call'](tool_name='qaqc_fast_session',args={'action':'write_record','expected_project':'Dự án TEST','record_id':'QAQC-PLUGIN-TEST-fixture'})
        self.assertEqual(verdict['action'],'approve')

    def test_workflow_transition_needs_separate_grant_and_changed_status(self):
        import hashlib
        record = 'QAQC-PLUGIN-TEST-FIXTURE-APPROVE'
        target = 'https://qaqc.newtecons.vn/#/editForm/fixture/42/1/project'
        before = snapshot('  uid=8_0 textbox "Mã hồ sơ" value="'+record+'"\n  uid=8_1 combobox "Trạng thái" value="Draft"\n  uid=8_2 button "Duyệt"\n').replace(URL,target)
        after = before.replace('value="Draft"','value="Approved"')
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[before,after,'Successfully reloaded the page.',after])
            old = host.get_config
            host.get_config = lambda key,default=None: 'Fixture Project' if key=='write_project' else ['approve'] if key=='workflow_transitions' else old(key,default)
            key = 'owned_'+hashlib.sha256(('Fixture Project\0'+record).encode()).hexdigest()
            host.state.set(key,{'project':'Fixture Project','user':'Fixture User','status':'created_verified','source_url':target})
            result = extended(host,dict(action='transition_record',page_id=1,expected_project='Fixture Project',record_id=record,operation_id='approve-fixture',transition='approve',control_label='Duyệt',expected_status='Approved',verify_fields={}))
            self.assertTrue(result['ok'],result)
            self.assertEqual(result['status'],'transition_verified_after_reload')

    def test_upload_requires_owned_record_and_persisted_attachment(self):
        import hashlib
        from openpyxl import Workbook
        record = 'QAQC-PLUGIN-TEST-FIXTURE-UPLOAD'
        target = 'https://qaqc.newtecons.vn/#/editForm/fixture/42/1/project'
        before = snapshot('  uid=8_0 textbox "Mã hồ sơ" value="'+record+'"\n  uid=8_1 button "Chọn tệp"\n  uid=8_2 button "Lưu"\n').replace(URL,target)
        after = before+'  uid=9_0 link "synthetic-upload.xlsx" url="https://qaqc.newtecons.vn/upload/synthetic-upload.xlsx"\n'
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'synthetic-upload.xlsx'
            Workbook().save(path)
            host = Host(root,[before,after,after,'Successfully reloaded the page.',after])
            old = host.get_config
            host.get_config = lambda key,default=None: 'Fixture Project' if key=='write_project' else old(key,default)
            key = 'owned_'+hashlib.sha256(('Fixture Project\0'+record).encode()).hexdigest()
            host.state.set(key,{'project':'Fixture Project','user':'Fixture User','status':'created_verified','source_url':target})
            result = extended(host,dict(action='upload_attachment',page_id=1,expected_project='Fixture Project',record_id=record,operation_id='upload-fixture',local_file=str(path),upload_label='Chọn tệp',submit_label='Lưu'))
            self.assertTrue(result['ok'],result)
            self.assertFalse(result['remote_byte_identity_verified'])
            self.assertEqual(host.calls[1][0],'upload_file')

    def test_edit_only_updates_plugin_owned_record(self):
        import hashlib
        record = 'QAQC-PLUGIN-TEST-FIXTURE-EDIT'
        target = 'https://qaqc.newtecons.vn/#/editForm/fixture/42/1/project'
        before = snapshot('  uid=8_0 textbox "Mã hồ sơ" value="'+record+'"\n  uid=8_1 textbox "Nội dung" value="Old"\n  uid=8_2 button "Lưu"\n  uid=8_3 combobox "Trạng thái" value="Draft"\n').replace(URL,target)
        after = before.replace('value="Old"','value="New"')
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[before,after,after,'Successfully reloaded the page.',after])
            old = host.get_config
            host.get_config = lambda key,default=None: 'Fixture Project' if key=='write_project' else old(key,default)
            key = 'owned_'+hashlib.sha256(('Fixture Project\0'+record).encode()).hexdigest()
            host.state.set(key,{'record_id':record,'project':'Fixture Project','user':'Fixture User','source_url':target,'status':'created_verified','fields':{'Mã hồ sơ':record}})
            result = extended(host,dict(action='write_record',page_id=1,expected_project='Fixture Project',record_id=record,operation_id='edit-fixture',intent='edit',fields={'Nội dung':'New'},submit_label='Lưu',expected_status='Draft'))
            self.assertTrue(result['ok'],result)
            self.assertEqual(result['intent'],'edit')

    def test_diagnostics_distinguish_closed_browser_without_input(self):
        import socket
        with socket.socket() as listener:
            listener.bind(('127.0.0.1',0))
            port = listener.getsockname()[1]
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[])
            old = host.get_config
            host.get_config = lambda key,default=None: port if key=='debug_port' else old(key,default)
            result = extended(host,dict(action='diagnose',page_id=1,expected_project='Fixture Project'))
            self.assertEqual(result['status'],'browser_closed',result)
            self.assertFalse(result['business_verified'])
            self.assertEqual(host.calls,[])

    def test_read_document_exposes_lossless_file_artifact(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'selected.xlsx'
            book = Workbook()
            book.save(path)
            host = Host(root,[snapshot()])
            result = extended(host,dict(action='read_document',page_id=1,expected_project='Fixture Project',local_file=str(path)))
            self.assertTrue(result['ok'],result)
            content = json.loads(Path(result['artifact']['path']).read_text(encoding='utf-8'))
            self.assertEqual(content['unit_count'],1)
            self.assertFalse(result['remote_source_link_verified'])

    def test_export_requires_new_valid_file_not_click_success(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'synthetic-export.xlsx'
            text = snapshot('  uid=8_0 button "Xuất Excel"\n')
            host = Host(root,[text,text,text])
            start = extended(host,dict(action='begin_transfer',page_id=1,expected_project='Fixture Project',transfer_kind='export_excel',control_label='Xuất Excel',local_file=str(path),operation_id='export-fixture'))
            self.assertEqual(start['status'],'triggered_waiting_file',start)
            book = Workbook()
            book.save(path)
            result = extended(host,dict(action='verify_transfer',page_id=1,expected_project='Fixture Project',operation_id='export-fixture'))
            self.assertTrue(result['ok'],result)
            self.assertTrue(result['file']['file_verified'])
            self.assertEqual(len(result['file']['sha256']),64)

    def test_pdf_reader_reads_each_text_page(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'synthetic.pdf'
            # A genuine generated PDF fixture; no test-only ReportLab in production.
            streams = [b'BT /F1 12 Tf 50 750 Td ('+text+b') Tj ET' for text in (b'Synthetic page one',b'Synthetic page two')]
            objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
                       b'<< /Type /Pages /Kids [3 0 R 5 0 R] /Count 2 >>',
                       b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << /Font << /F1 7 0 R >> >> /Contents 4 0 R >>',
                       b'<< /Length '+str(len(streams[0])).encode()+b' >>\nstream\n'+streams[0]+b'\nendstream',
                       b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << /Font << /F1 7 0 R >> >> /Contents 6 0 R >>',
                       b'<< /Length '+str(len(streams[1])).encode()+b' >>\nstream\n'+streams[1]+b'\nendstream',
                       b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
            payload,offsets = b'%PDF-1.4\n',[]
            for index,obj in enumerate(objects,1):
                offsets.append(len(payload))
                payload += str(index).encode()+b' 0 obj\n'+obj+b'\nendobj\n'
            xref = len(payload)
            payload += b'xref\n0 '+str(len(objects)+1).encode()+b'\n0000000000 65535 f \n'
            payload += b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets)
            payload += b'trailer\n<< /Size '+str(len(objects)+1).encode()+b' /Root 1 0 R >>\nstartxref\n'+str(xref).encode()+b'\n%%EOF\n'
            path.write_bytes(payload)
            reader = cast(Any,plugin.read_document_file)
            result = cast(Any,reader(path))
            self.assertTrue(result['ok'],result)
            self.assertEqual(result['unit_count'],2)
            self.assertIn('Synthetic page two',result['units'][1]['text'])

    def test_native_step_never_replays_unverified_input(self):
        import hashlib
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[])
            record = 'QAQC-PLUGIN-TEST-FIXTURE-NATIVE'
            old = host.get_config
            host.get_config = lambda key,default=None: 'Fixture Project' if key=='write_project' else old(key,default)
            image = Path(root)/'capture.png'
            image.write_bytes(b'synthetic capture')
            digest = hashlib.sha256(image.read_bytes()).hexdigest()
            host.state.set('native_handoff',{'token':'fixture-token','session_id':'fixture-session','project':'Fixture Project','user':'Fixture User','record_id':record,'source_url':URL,'sha256':digest,'screenshot_path':str(image),'captured_at':__import__('time').time(),'consumed':False,'elements':{'1':'Canvas','2':'Canvas'},'pid':1,'window_id':2})
            key = 'owned_'+hashlib.sha256(('Fixture Project\0'+record).encode()).hexdigest()
            host.state.set(key,{'project':'Fixture Project','user':'Fixture User','source_url':URL,'status':'created_verified'})
            calls = []
            def dispatch(name,args,**kw):
                if args['action']=='capture':
                    return json.dumps({'screenshot_path':str(image),'text_summary':'#1 Canvas\n#2 Canvas'})
                calls.append(args)
                return json.dumps({'ok':True,'effect':'unverifiable'})
            cast(Any,host).dispatch_tool = dispatch
            args = dict(action='native_step',page_id=1,expected_project='Fixture Project',record_id=record,operation_id='native-fixture',handoff_token='fixture-token',reviewed_capture_sha256=digest,purpose='canvas',native_action={'action':'drag','from_element':1,'to_element':2})
            first = extended(host,args)
            self.assertEqual(first['status'],'native_effect_unverified')
            second = extended(host,args)
            self.assertFalse(second['ok'])
            self.assertEqual(len(calls),1)

    def test_document_reader_preserves_all_workbook_units(self):
        from openpyxl import Workbook
        reader = cast(Any,getattr(plugin,'read_document_file',None))
        self.assertTrue(callable(reader),'Missing bounded document reader')
        assert callable(reader)
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'synthetic.xlsx'
            book = Workbook()
            sheet = book.active
            assert sheet is not None
            sheet.title = 'Visible'
            for i in range(130):
                sheet.append([f'R{i}',i])
            hidden = book.create_sheet('Hidden')
            hidden.sheet_state = 'hidden'
            hidden['A1'] = '=SUM(Visible!B1:B130)'
            book.save(path)
            result = reader(path)
            self.assertTrue(result['ok'],result)
            self.assertEqual(len(result['units']),2)
            self.assertEqual(len(result['units'][0]['rows']),130)
            self.assertEqual(result['units'][1]['rows'][0][0],'=SUM(Visible!B1:B130)')
            self.assertFalse(result['formulas_calculated'])

    def test_create_record_verifies_persistence_after_reload(self):
        record = 'QAQC-PLUGIN-TEST-FIXTURE-001'
        target = 'https://qaqc.newtecons.vn/#/editForm/fixture/0/1/project'
        listing = snapshot('  uid=8_0 button "Thêm mới"\n')
        form = snapshot('  uid=9_0 textbox "Mã hồ sơ" value=""\n  uid=9_1 textbox "Nội dung" value=""\n  uid=9_2 button "Lưu"\n').replace(URL,target)
        filled = form.replace('value=""','value="'+record+'"',1).replace('value=""','value="Synthetic content"',1)
        saved = filled.replace('/fixture/0/','/fixture/42/')+'  uid=10_0 combobox "Trạng thái" value="Draft"\n'
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[listing,form,filled,saved,'Successfully reloaded the page.',saved])
            old = host.get_config
            host.get_config = lambda key,default=None: 'Fixture Project' if key=='write_project' else old(key,default)
            result = extended(host,dict(action='write_record',page_id=1,expected_project='Fixture Project',intent='create',record_id=record,operation_id='create-fixture-001',fields={'Mã hồ sơ':record,'Nội dung':'Synthetic content'},create_label='Thêm mới',submit_label='Lưu',expected_status='Draft'))
            self.assertTrue(result['ok'],result)
            self.assertEqual(result['status'],'write_verified_after_reload')
            self.assertEqual([x[0] for x in host.calls],['take_snapshot','click','fill_form','click','navigate_page','take_snapshot'])

    def test_native_capture_is_bound_to_verified_qaqc_context(self):
        with tempfile.TemporaryDirectory() as root:
            png = Path(root)/'capture.png'
            png.write_bytes(b'fixture screenshot')
            host = Host(root,[snapshot()])
            payload = {'ok':True,'screenshot_path':str(png),'text_summary':URL+' Fixture Project Fixture User','_multimodal':True,'content':[{'type':'text','text':'fixture AX tree'}]}
            host.dispatch_tool = lambda name,args,**kwargs: json.dumps(payload)
            result = extended(host,dict(action='native_capture',page_id=1,expected_project='Fixture Project',pid=101,window_id=202))
            self.assertTrue(result['ok'],result)
            self.assertTrue(result['_multimodal'])
            self.assertIn('handoff',result)
            self.assertEqual(host.state.get('native_handoff')['project'],'Fixture Project')

    def test_open_record_reads_identity_bound_detail_tabs(self):
        target = 'https://qaqc.newtecons.vn/#/viewForm/fixture/42/1/project'
        listing = snapshot('  uid=8_0 row\n    uid=8_1 cell "R001"\n    uid=8_2 link "Xem" url="'+target+'"\n')
        detail = snapshot('  uid=9_0 heading "R001"\n  uid=9_1 tab "Lịch sử" selected=false\n  uid=9_2 StaticText "Trạng thái: Draft"\n').replace(URL,target)
        history = detail.replace('selected=false','selected=true')+'  uid=10_0 StaticText "History event"\n'
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[listing,detail,history])
            result = extended(host,dict(action='open_record',page_id=1,expected_project='Fixture Project',record_id='R001',detail_label='Xem',tabs=['Lịch sử']))
            self.assertTrue(result['ok'],result)
            self.assertTrue(result['record_identity_verified'])
            value = json.loads(Path(result['artifact']['path']).read_text(encoding='utf-8'))
            self.assertIn('History event',value['tabs'][1]['data']['visible_text'])

    def test_read_all_reconciles_identity_across_pages(self):
        def page(codes,disabled):
            body = '  uid=7_0 StaticText "Total: 3"\n  uid=8_0 table busy=false\n'
            body += ''.join(f'    uid=9_{i} row\n      uid=10_{i} cell "{code}"\n' for i,code in enumerate(codes))
            return snapshot(body+'  uid=11_0 button "Next" disabled='+str(disabled).lower()+'\n')
        first,second = page(['R1','R2'],False),page(['R2','R3'],True)
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[first,first,second,second])
            result = extended(host,dict(action='read_all',page_id=1,expected_project='Fixture Project',next_label='Next',key_column=0,max_pages=5))
            self.assertTrue(result.get('inventory_complete'),result)
            self.assertEqual(result['unique_records'],3)
            self.assertEqual(result['duplicates'],1)
            manifest = json.loads(Path(result['artifact']['path']).read_text(encoding='utf-8'))
            self.assertEqual(len(manifest['batches']),2)
            self.assertEqual(len(manifest['records']),3)

    def test_query_verifies_filters_before_reading_results(self):
        before = snapshot('  uid=8_0 textbox "Từ khóa" value=""\n  uid=8_1 combobox "Loại hồ sơ" value="Tất cả"\n  uid=8_2 button "Tìm kiếm"\n')
        filled = before.replace('value=""','value="R001"').replace('value="Tất cả"','value="Nghiệm thu"')
        loaded = filled + '  uid=9_0 table busy=false\n    uid=9_1 row\n      uid=9_2 cell "R001"\n'
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[before,filled,loaded,loaded])
            result = extended(host,dict(action='query',page_id=1,expected_project='Fixture Project',fields={'Từ khóa':'R001','Loại hồ sơ':'Nghiệm thu'},submit_label='Tìm kiếm'))
            self.assertTrue(result['ok'],result)
            self.assertEqual(result['rows'],1)
            self.assertEqual([x[0] for x in host.calls],['take_snapshot','fill_form','click','take_snapshot'])

    def test_set_context_verifies_exact_target_context(self):
        before = snapshot('  uid=8_0 link "Target Project" url="'+URL+'"\n')
        after = snapshot('  uid=9_0 tab "CƠ ĐIỆN" selected=false\n',project='Target Project')
        selected = after.replace('selected=false','selected=true')
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[before,after,selected])
            result = extended(host,dict(action='set_context',page_id=1,from_project='Fixture Project',expected_project='Target Project',expected_discipline='CƠ ĐIỆN'))
            self.assertTrue(result['ok'],result)
            self.assertTrue(result['discipline_verified'])
            self.assertEqual(result['source']['project'],'Target Project')

    def test_navigation_uses_observed_route_and_verifies_destination(self):
        target = 'https://qaqc.newtecons.vn/#/danh-sach-nhat-ky-thi-cong/1/fixture'
        before = snapshot(f'  uid=8_0 link "Nhật ký thi công" url="{target}"\n')
        after = snapshot('  uid=8_0 heading "NHẬT KÝ THI CÔNG"\n').replace(URL,target)
        with tempfile.TemporaryDirectory() as root:
            host = Host(root,[before,after])
            result = extended(host,dict(action='navigate',page_id=1,expected_project='Fixture Project',menu_path=['Nhật ký thi công'],destination_heading='NHẬT KÝ THI CÔNG'))
            self.assertTrue(result['ok'],result)
            self.assertEqual(result['source']['url'],target)
            self.assertEqual([x[0] for x in host.calls],['take_snapshot','click'])

    def test_read_view_saves_lossless_verified_data(self):
        with tempfile.TemporaryDirectory() as root:
            body = '  uid=2_0 table busy=false\n'
            body += ''.join(f'    uid=3_{i} row\n      uid=4_{i} cell "R{i}"\n' for i in range(120))
            text = snapshot(body,discipline='  uid=9_0 tab "XÂY DỰNG" selected=true\n')
            host = Host(root,[text,text])
            args = dict(action='read_view',page_id=1,expected_project='Fixture Project',expected_discipline='XÂY DỰNG')
            result = extended(host,args)
            self.assertTrue(result['ok'],result)
            self.assertTrue(result['data_ready_verified'])
            self.assertTrue(result['discipline_verified'])
            data = json.loads(Path(result['artifact']['path']).read_text(encoding='utf-8'))
            self.assertEqual(len(data['list_data']['rows']),120)
            self.assertEqual(result['coverage'],'current_page_only')

    def test_lossless_visible_text_rows_and_links(self):
        body = ''.join(f'  uid=2_{i} StaticText "Text {i}"\n' for i in range(280))
        body += ''.join(f'  uid=3_{i} row\n    uid=4_{i} cell "Record {i}"\n' for i in range(120))
        body += ''.join(f'  uid=5_{i} link "Document {i}" url="https://qaqc.newtecons.vn/#/AlbumThiCong/{i}/x"\n' for i in range(110))
        data = plugin.list_data(plugin.parse_snapshot(snapshot(body)))
        self.assertEqual(len(data['visible_text']), 280)
        self.assertEqual(len(data['rows']), 120)
        self.assertEqual(len(data['source_links']), 110)
        self.assertFalse(data['visible_text_truncated'])
        self.assertFalse(data['rows_truncated'])
        self.assertFalse(data['links_truncated'])

    def test_readiness_requires_stable_data_and_completion_signal(self):
        fn = cast(Any,getattr(plugin, 'readiness', None))
        self.assertTrue(callable(fn), 'Missing readiness verifier')
        assert callable(fn)
        shell = plugin.parse_snapshot(snapshot('  uid=2_0 heading "LIST"\n'))
        loaded = plugin.parse_snapshot(snapshot('  uid=2_0 table busy=false\n    uid=2_1 row\n      uid=2_2 cell "R001"\n'))
        self.assertFalse(fn(shell,shell)['data_ready_verified'])
        self.assertFalse(fn(shell,loaded)['data_ready_verified'])
        self.assertTrue(fn(loaded,loaded)['data_ready_verified'])
        busy = plugin.parse_snapshot(snapshot('  uid=2_0 table busy=true\n    uid=2_1 row\n      uid=2_2 cell "R001"\n'))
        self.assertFalse(fn(busy,busy)['data_ready_verified'])

    def test_discipline_requires_selected_evidence(self):
        fn = getattr(plugin, 'selected_discipline', None)
        self.assertTrue(callable(fn), 'Missing selected discipline verifier')
        assert callable(fn)
        for label in ('XÂY DỰNG','CƠ ĐIỆN'):
            nodes = plugin.parse_snapshot(snapshot(discipline=f'  uid=9_0 tab "{label}" selected=true\n'))
            self.assertEqual(fn(nodes), label)
        self.assertIsNone(fn(plugin.parse_snapshot(snapshot())))

    def test_top_page_after_iframe_is_kept(self):
        text = snapshot('  uid=2_0 StaticText "Before"\n'
                        '  uid=2_1 Iframe\n'
                        '    uid=2_2 RootWebArea "Frame" url="https://example.com/"\n'
                        '      uid=2_3 StaticText "Foreign"\n'
                        '  uid=2_4 StaticText "After"\n')
        data = plugin.list_data(plugin.parse_snapshot(text))
        self.assertEqual(data['visible_text'], ['Before','After'])

if __name__ == '__main__':
    unittest.main()
